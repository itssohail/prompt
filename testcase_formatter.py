"""
Test Case Formatter
-------------------
Paste the table your Rovo agent gives you and get a colorful, formatted .xlsx file.

What it does
  * Adds the column header row (Rovo doesn't give one)
  * Groups cases under STORY and CATEGORY heading rows (adds them itself)
  * Turns " || " into real line breaks inside cells (Test Steps / Expected Result)
  * Colors Priority, Status and Test Type; adds dropdowns for Status / Priority / Type
  * Builds a Summary sheet (formulas) with counts by type, priority, status and story
  * Reports rows that don't have exactly 16 columns

Setup (one time):   pip install openpyxl
Run with window:    python testcase_formatter.py
Run from a file:    python testcase_formatter.py input.txt [output.xlsx]
"""

import math
import os
import re
import sys
from collections import OrderedDict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col_letter
from openpyxl.worksheet.datavalidation import DataValidation

# --------------------------------------------------------------------------
# Configuration
# --------------------------------------------------------------------------
COLUMNS = [
    ("Test Case ID", 18), ("Story Key", 13), ("Story Summary", 28), ("AC Reference", 12),
    ("Module / Feature", 20), ("Test Scenario", 30), ("Test Case Title", 38), ("Test Type", 14),
    ("Priority", 11), ("Preconditions", 30), ("Test Data", 28), ("Test Steps", 52),
    ("Expected Result", 52), ("Actual Result", 26), ("Status", 14), ("Comments / Notes", 30),
]
NCOLS = len(COLUMNS)
IDX_KEY, IDX_TYPE, IDX_PRIORITY, IDX_STATUS = 1, 7, 8, 14
IDX_SUMMARY = 2

CATEGORY_ORDER = ["Functional", "UI", "Negative", "Boundary", "Integration", "Security", "Regression"]
TYPE_LOOKUP = {t.lower(): t for t in CATEGORY_ORDER}
PRIORITY_LOOKUP = {"high": "High", "medium": "Medium", "low": "Low"}
STATUS_LOOKUP = {"not executed": "Not Executed", "pass": "Pass", "fail": "Fail", "blocked": "Blocked"}

FONT = "Arial"
C_HEADER = "1F4E78"
C_STORY = "1F4E78"
C_CATEGORY = "9DC3E6"
C_BAND = "F2F7FB"
PRIORITY_COLORS = {"High": "F8CBAD", "Medium": "FFE699", "Low": "C6E0B4"}
STATUS_COLORS = {"Not Executed": "D9D9D9", "Pass": "A9D08E", "Fail": "FF7C80", "Blocked": "F4B183"}
TYPE_COLORS = {
    "Functional": "DDEBF7", "UI": "E4DFEC", "Negative": "FCE4D6", "Boundary": "FFF2CC",
    "Integration": "E2EFDA", "Security": "FFD7D7", "Regression": "EDEDED",
}

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color, bgColor=hex_color)


# --------------------------------------------------------------------------
# Parsing
# --------------------------------------------------------------------------
def _is_heading(text):
    t = text.strip().upper()
    return t.startswith("STORY:") or t.startswith("CATEGORY:")


def parse(text):
    """Return (cases, warnings). cases = list of 16-item lists."""
    warnings = []
    cases = []
    text = text.replace("\ufeff", "")
    lines = text.splitlines()

    for n, raw in enumerate(lines, 1):
        line = raw.rstrip("\r\n")
        stripped = line.strip()
        if not stripped or stripped.startswith("```"):
            continue
        if _is_heading(stripped) and "\t" not in line:
            continue  # heading line from Rovo; the tool rebuilds headings itself

        if "\t" in line:
            cells = line.split("\t")
        elif stripped.count("|") >= 3:
            # markdown-style table fallback; protect the ' || ' line-break marker
            if re.fullmatch(r"[\s|:\-]+", stripped):
                continue
            protected = re.sub(r"\s*\|\|\s*", "\x00", stripped)
            protected = protected.strip("|")
            cells = [c.replace("\x00", " || ") for c in protected.split("|")]
        else:
            warnings.append(f"Line {n}: could not split into columns (no tabs found) - skipped.")
            continue

        cells = [c.strip() for c in cells]
        # skip a header row if Rovo ever includes one
        if cells and cells[0].lower() == "test case id":
            continue
        # skip heading rows that arrived with trailing tabs
        if cells and _is_heading(cells[0]) and not any(cells[1:]):
            continue

        if len(cells) < NCOLS:
            warnings.append(f"Line {n}: has {len(cells)} columns, expected {NCOLS} - padded with blanks. "
                            f"({cells[0][:30]})")
            cells += [""] * (NCOLS - len(cells))
        elif len(cells) > NCOLS:
            warnings.append(f"Line {n}: has {len(cells)} columns, expected {NCOLS} - extra text merged into the last "
                            f"column. ({cells[0][:30]})")
            cells = cells[:NCOLS - 1] + [" ".join(cells[NCOLS - 1:])]

        cases.append(cells)

    return cases, warnings


def clean_cell(value):
    v = value.replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
    v = re.sub(r"\s*\|\|\s*", "\n", v)
    return v.strip()


def normalise(cases):
    """Tidy values: line breaks, type / priority / status spelling, default status."""
    out = []
    for row in cases:
        row = [clean_cell(c) for c in row]
        row[IDX_TYPE] = TYPE_LOOKUP.get(row[IDX_TYPE].lower(), row[IDX_TYPE])
        row[IDX_PRIORITY] = PRIORITY_LOOKUP.get(row[IDX_PRIORITY].lower(), row[IDX_PRIORITY])
        st = STATUS_LOOKUP.get(row[IDX_STATUS].lower(), row[IDX_STATUS])
        row[IDX_STATUS] = st or "Not Executed"
        out.append(row)
    return out


def group(cases):
    """OrderedDict: story_key -> {'summary': str, 'cats': OrderedDict(category -> [rows])}"""
    stories = OrderedDict()
    for row in cases:
        key = row[IDX_KEY] or "UNKNOWN"
        s = stories.setdefault(key, {"summary": row[IDX_SUMMARY], "cats": {}})
        if not s["summary"] and row[IDX_SUMMARY]:
            s["summary"] = row[IDX_SUMMARY]
        cat = row[IDX_TYPE] if row[IDX_TYPE] else "Other"
        s["cats"].setdefault(cat, []).append(row)
    for s in stories.values():
        ordered = OrderedDict()
        for c in CATEGORY_ORDER:
            if c in s["cats"]:
                ordered[c] = s["cats"][c]
        for c, rows in s["cats"].items():  # anything unexpected goes last
            if c not in ordered:
                ordered[c] = rows
        s["cats"] = ordered
    return stories


# --------------------------------------------------------------------------
# Excel output
# --------------------------------------------------------------------------
def _row_height(row_values):
    lines = 1
    for i, val in enumerate(row_values):
        if not val:
            continue
        width = COLUMNS[i][1]
        per_line = max(1, int(width * 1.05))
        n = sum(max(1, math.ceil(len(part) / per_line)) for part in str(val).split("\n"))
        lines = max(lines, n)
    return min(409, max(18, 13 * lines + 4))


def build_workbook(cases, out_path):
    stories = group(cases)
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    # header row
    for i, (name, width) in enumerate(COLUMNS, 1):
        c = ws.cell(1, i, name)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
        c.fill = fill(C_HEADER)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[col_letter(i)].width = width
    ws.row_dimensions[1].height = 32

    r = 2
    case_rows = 0
    band = 0
    for key, s in stories.items():
        total = sum(len(v) for v in s["cats"].values())
        title = f"STORY: {key}" + (f" - {s['summary']}" if s["summary"] else "") + f"   ({total} test cases)"
        _heading_row(ws, r, title, C_STORY, "FFFFFF", 12)
        r += 1
        for cat, rows in s["cats"].items():
            _heading_row(ws, r, f"CATEGORY: {cat} Test Cases ({len(rows)})", C_CATEGORY, "000000", 11)
            r += 1
            for row in rows:
                for i, val in enumerate(row, 1):
                    c = ws.cell(r, i, val)
                    c.font = Font(name=FONT, size=10)
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                    c.border = BORDER
                    if band % 2 == 1:
                        c.fill = fill(C_BAND)
                for i in (IDX_TYPE, IDX_PRIORITY, IDX_STATUS):
                    ws.cell(r, i + 1).alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
                ws.row_dimensions[r].height = _row_height(row)
                r += 1
                band += 1
                case_rows += 1
    last = r - 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{col_letter(NCOLS)}{max(last, 2)}"

    # dropdowns
    if last >= 2:
        def dv(options, col):
            d = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
            ws.add_data_validation(d)
            d.add(f"{col}2:{col}{last}")
        dv(CATEGORY_ORDER, col_letter(IDX_TYPE + 1))
        dv(["High", "Medium", "Low"], col_letter(IDX_PRIORITY + 1))
        dv(list(STATUS_COLORS), col_letter(IDX_STATUS + 1))

        # conditional colors (recolor automatically when a dropdown value changes)
        def cf(col, colors):
            L = col_letter(col + 1)
            for value, color in colors.items():
                ws.conditional_formatting.add(
                    f"{L}2:{L}{last}",
                    FormulaRule(formula=[f'{L}2="{value}"'], fill=fill(color),
                                font=Font(name=FONT, bold=True, size=10), stopIfTrue=True))
        cf(IDX_PRIORITY, PRIORITY_COLORS)
        cf(IDX_STATUS, STATUS_COLORS)
        cf(IDX_TYPE, TYPE_COLORS)

    _summary_sheet(wb, stories, last)
    wb.save(out_path)
    return case_rows, len(stories)


def _heading_row(ws, r, text, color, font_color, size):
    for i in range(1, NCOLS + 1):
        c = ws.cell(r, i)
        c.fill = fill(color)
        c.border = Border(top=THIN, bottom=THIN)
    a = ws.cell(r, 1, text)
    a.font = Font(name=FONT, bold=True, color=font_color, size=size)
    a.alignment = Alignment(wrap_text=False, vertical="center")
    ws.row_dimensions[r].height = 22


def _summary_sheet(wb, stories, last):
    sm = wb.create_sheet("Summary")
    sm.sheet_properties.tabColor = "70AD47"
    sm.column_dimensions["A"].width = 34
    sm.column_dimensions["B"].width = 14
    last = max(last, 2)
    rng = lambda letter: f"'Test Cases'!${letter}$2:${letter}${last}"

    sm["A1"] = "Test Design Summary"
    sm["A1"].font = Font(name=FONT, bold=True, size=14, color=C_HEADER)

    row = [3]

    def section(title):
        for c in (1, 2):
            x = sm.cell(row[0], c)
            x.fill = fill("2E75B6")
            x.font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
            x.border = BORDER
        sm.cell(row[0], 1, title)
        row[0] += 1

    def line(label, formula, fmt="0", color=None):
        a = sm.cell(row[0], 1, label)
        b = sm.cell(row[0], 2, formula)
        for x in (a, b):
            x.font = Font(name=FONT, size=10)
            x.border = BORDER
        b.number_format = fmt
        b.alignment = Alignment(horizontal="center")
        if color:
            a.fill = fill(color)
        row[0] += 1
        return row[0] - 1

    section("Overview")
    line("Stories", f'=COUNTIF({rng("A")},"STORY:*")')
    total_row = line("Total test cases", f'=COUNTIF({rng("B")},"?*")')
    row[0] += 1
    section("By Test Type")
    for t in CATEGORY_ORDER:
        line(t, f'=COUNTIF({rng("H")},"{t}")', color=TYPE_COLORS[t])
    row[0] += 1
    section("By Priority")
    for p in ("High", "Medium", "Low"):
        line(p, f'=COUNTIF({rng("I")},"{p}")', color=PRIORITY_COLORS[p])
    row[0] += 1
    section("Execution Status")
    pass_row = None
    for s_ in STATUS_COLORS:
        rr = line(s_, f'=COUNTIF({rng("O")},"{s_}")', color=STATUS_COLORS[s_])
        if s_ == "Pass":
            pass_row = rr
    line("Pass %", f"=IFERROR(B{pass_row}/B{total_row},0)", "0.0%")
    row[0] += 1
    section("Test Cases per Story")
    for key in stories:
        line(key, f'=COUNTIF({rng("B")},"{key}")')


# --------------------------------------------------------------------------
# Entry points
# --------------------------------------------------------------------------
def convert(text, out_path):
    cases, warnings = parse(text)
    if not cases:
        raise ValueError("No test case rows found. Make sure you pasted the table Rovo generated.")
    cases = normalise(cases)
    n_cases, n_stories = build_workbook(cases, out_path)
    return n_cases, n_stories, warnings


def default_output_name():
    return f"Test_Cases_{datetime.now():%Y%m%d_%H%M}.xlsx"


def run_gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext

    root = tk.Tk()
    root.title("Test Case Formatter")
    root.geometry("900x600")

    tk.Label(root, text="Paste the table from your Rovo agent below, then click Generate Excel.",
             font=("Arial", 11)).pack(anchor="w", padx=10, pady=(10, 4))
    box = scrolledtext.ScrolledText(root, wrap="none", font=("Consolas", 9))
    box.pack(fill="both", expand=True, padx=10, pady=4)

    def paste():
        try:
            box.delete("1.0", "end")
            box.insert("1.0", root.clipboard_get())
        except tk.TclError:
            messagebox.showwarning("Clipboard", "Clipboard is empty.")

    def generate():
        text = box.get("1.0", "end")
        if not text.strip():
            messagebox.showwarning("Nothing to convert", "Paste the Rovo output first.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile=default_output_name(),
                                            filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        try:
            n_cases, n_stories, warnings = convert(text, path)
        except Exception as e:  # show any problem instead of crashing
            messagebox.showerror("Could not create file", str(e))
            return
        msg = f"Created {os.path.basename(path)}\n{n_cases} test cases across {n_stories} story(ies)."
        if warnings:
            msg += "\n\nWarnings:\n" + "\n".join(warnings[:15])
            if len(warnings) > 15:
                msg += f"\n...and {len(warnings) - 15} more."
        msg += "\n\nOpen the file now?"
        if messagebox.askyesno("Done", msg):
            try:
                if sys.platform.startswith("win"):
                    os.startfile(path)
                elif sys.platform == "darwin":
                    os.system(f'open "{path}"')
                else:
                    os.system(f'xdg-open "{path}"')
            except Exception:
                pass

    bar = tk.Frame(root)
    bar.pack(pady=8)
    tk.Button(bar, text="Paste from clipboard", command=paste, width=20).pack(side="left", padx=5)
    tk.Button(bar, text="Generate Excel", command=generate, width=20, bg="#1F4E78", fg="white").pack(side="left", padx=5)
    tk.Button(bar, text="Clear", command=lambda: box.delete("1.0", "end"), width=10).pack(side="left", padx=5)
    root.mainloop()


def main():
    if len(sys.argv) >= 2:
        src = sys.argv[1]
        out = sys.argv[2] if len(sys.argv) >= 3 else default_output_name()
        with open(src, "r", encoding="utf-8-sig") as f:
            text = f.read()
        n_cases, n_stories, warnings = convert(text, out)
        print(f"Created {out}: {n_cases} test cases, {n_stories} story(ies).")
        for w in warnings:
            print("WARNING:", w)
    else:
        run_gui()


if __name__ == "__main__":
    main()
